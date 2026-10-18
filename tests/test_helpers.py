from datetime import date
from decimal import Decimal as D

import pytest

from app.formatting import format_inr, format_pct
from app.forms import (
    clean_text, parse_choice, parse_date, parse_money, parse_optional_int, parse_positive_int, parse_price,
    parse_quantity, parse_rate,
)
from app.services.charts import line_chart
from app.services.errors import ValidationError
from app.services.exports import to_csv, to_xlsx
from app.services.periods import (
    fy_bounds, fy_label, fy_year_of, month_bounds, resolve_range, shift_month,
)


def test_inr_uses_indian_grouping():
    assert format_inr(D("1234567.5")) == "₹12,34,567.50"
    assert format_inr(D("12345678.9")) == "₹1,23,45,678.90"
    assert format_inr(D("999")) == "₹999.00"
    assert format_inr(D("1000")) == "₹1,000.00"
    assert format_inr(D("-100000")) == "-₹1,00,000.00"
    assert format_inr(D("1234.5"), 0) == "₹1,235"
    assert format_inr(None) == "-"


def test_pct():
    assert format_pct(D("66.04")) == "66.0%"
    assert format_pct(None) == "n/a"


def test_parse_money():
    assert parse_money("₹ 1,25,000.456", "Amount") == D("125000.46")
    assert parse_money("-5", "x", allow_negative=True) == D("-5.00")
    assert parse_money("0", "x", allow_zero=True) == D("0.00")
    for bad in ("", "abc", "-5", "0", "NaN", "Infinity", "1e20"):
        with pytest.raises(ValidationError):
            parse_money(bad, "Amount")


def test_parse_date_and_choice_and_text():
    assert parse_date("2026-10-18", "Date") == date(2026, 10, 18)
    with pytest.raises(ValidationError):
        parse_date("18/10/2026", "Date")
    with pytest.raises(ValidationError):
        parse_choice("nope", "Type", ["a", "b"])
    assert clean_text("  hi  ", "x") == "hi"
    assert clean_text("", "x") is None
    with pytest.raises(ValidationError):
        clean_text("", "Name", required=True)
    with pytest.raises(ValidationError):
        clean_text("x" * 11, "Name", 10)
    assert parse_optional_int("", "x") is None
    with pytest.raises(ValidationError, match="Choose an account"):
        parse_optional_int("", "Account", required=True)
    with pytest.raises(ValidationError, match="Choose a category"):
        parse_optional_int("", "Category", required=True)
    with pytest.raises(ValidationError):
        parse_optional_int("abc", "x")
    assert parse_rate("8.5%", "Rate") == D("8.50")
    with pytest.raises(ValidationError):
        parse_rate("101", "Rate")


def test_financial_year_boundaries():
    assert fy_year_of(date(2026, 3, 31)) == 2025
    assert fy_year_of(date(2026, 4, 1)) == 2026
    assert fy_bounds(2026) == (date(2026, 4, 1), date(2027, 3, 31))
    assert fy_label(2026) == "FY 2026-27"
    assert fy_label(2099) == "FY 2099-00"


def test_month_helpers():
    assert month_bounds(date(2024, 2, 10)) == (date(2024, 2, 1), date(2024, 2, 29))
    assert shift_month(date(2026, 1, 15), -1) == date(2025, 12, 1)
    assert shift_month(date(2026, 11, 5), 3) == date(2027, 2, 1)


def test_resolve_range():
    today = date(2026, 10, 18)
    assert resolve_range("this_month", today) == (date(2026, 10, 1), date(2026, 10, 31), "this_month")
    assert resolve_range("last_month", date(2026, 1, 10)) == (date(2025, 12, 1), date(2025, 12, 31), "last_month")
    assert resolve_range("this_fy", today) == (date(2026, 4, 1), date(2027, 3, 31), "this_fy")
    assert resolve_range("last_fy", today) == (date(2025, 4, 1), date(2026, 3, 31), "last_fy")
    assert resolve_range("all", today) == (None, None, "all")
    assert resolve_range("custom", today, date(2026, 1, 1), None) == (date(2026, 1, 1), None, "custom")
    assert resolve_range("bogus", today)[2] == "this_month"


def test_csv_export_neutralises_formulas():
    out = to_csv(["a", "b"], [["=1+1", "5.00"], [None, "ok"], ["+cmd", "-x"]]).decode("utf-8-sig")
    lines = out.splitlines()
    assert lines[1] == "'=1+1,5.00"
    assert lines[2] == ",ok"
    assert lines[3] == "'+cmd,'-x"


def test_xlsx_export_roundtrip():
    from io import BytesIO

    from openpyxl import load_workbook

    data = to_xlsx(["Date", "Amount", "Note"], [[date(2026, 10, 1), D("12.50"), "=evil"]], "Test")
    ws = load_workbook(BytesIO(data)).active
    assert ws["A1"].value == "Date"
    assert float(ws["B2"].value) == 12.5
    assert ws["C2"].value == "'=evil"


def test_line_chart_geometry():
    assert line_chart([]) is None
    assert len(line_chart([1, 2, 3])["dots"]) == 3
    assert line_chart([5])["dots"]
    flat = line_chart([4, 4, 4])
    assert len({y for _x, y in flat["dots"]}) == 1


def test_quantity_price_and_ratio_parsing():
    assert parse_quantity("1,250") == D("1250.000000")
    assert parse_quantity("12.3456789") == D("12.345679")
    assert parse_price("₹ 2,450.755") == D("2450.7550")
    assert parse_positive_int("5", "Ratio") == 5
    for fn, bad in ((parse_quantity, "0"), (parse_quantity, "-1"), (parse_quantity, "x"), (parse_price, "0"),
                    (parse_price, ""), (parse_positive_int, "0"), (parse_positive_int, "1.5"), (parse_positive_int, "")):
        with pytest.raises(ValidationError):
            fn(bad) if fn is not parse_positive_int else fn(bad, "Ratio")
